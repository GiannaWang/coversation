from openpyxl import load_workbook, Workbook

# ===================== Config =====================
# Use a raw string or double backslashes for Windows paths.
EXCEL_FILE_PATH = r"F:\Desktop\your_file.xlsx"
SHEET_NAME = "Sheet1"

# ===================== Target rows =====================
target_rows1 = [
    569, 568, 567, 566, 563, 562, 559, 557, 556, 554, 552, 551, 550, 549, 548, 547, 546, 542, 541, 540,
    539, 537, 535, 533, 532, 530, 529, 528, 527, 526, 525, 524, 522, 521, 520, 518, 516, 515, 512, 510,
    509, 508, 504, 501, 498, 497, 496, 495, 493, 492, 491, 490, 489, 488, 487, 486, 484, 483, 482, 480,
    479, 478, 476, 475, 471, 469, 468, 467, 465, 464, 463, 462, 460, 458, 455, 452, 449, 448, 441, 439,
    438, 437, 436, 435, 433, 432, 430, 427, 425, 424, 422, 421, 420, 419, 416, 415, 414, 411, 407, 406,
    405, 404, 403, 401, 400, 399, 394, 393, 392, 391, 390, 388, 385, 384, 381, 379, 376, 375, 374, 373,
    372, 370, 369, 368, 367, 366, 365, 364, 363, 361, 360, 358, 356, 355, 353, 352, 350, 349, 348, 345,
    344, 343, 342, 341, 340, 339, 338, 337, 336, 329, 328, 327, 326, 325, 323, 322, 320, 317, 314, 313,
    311, 309, 308, 301, 299, 298, 297, 296, 295, 291, 290, 287, 286, 284, 283, 282, 281, 280,
]

target_rows2 = [
    565, 553, 544, 543, 517, 513, 511, 503, 502, 485, 481, 473, 472, 470, 466, 461, 456, 454, 453, 450,
    446, 445, 442, 434, 431, 423, 417, 413, 412, 409, 408, 389, 382, 380, 347, 335, 319, 316, 312, 310,
    307, 303, 300, 288,
]


def extract_excel_f_column(file_path: str, sheet_name: str, rows_list: list[int]) -> list:
    """Extract F column values from specified rows."""
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    f_column_data = []

    for row_num in rows_list:
        cell_value = ws[f"F{row_num}"].value
        f_column_data.append(cell_value)

    wb.close()
    return f_column_data


data1 = extract_excel_f_column(EXCEL_FILE_PATH, SHEET_NAME, target_rows1)
data2 = extract_excel_f_column(EXCEL_FILE_PATH, SHEET_NAME, target_rows2)

print("=== Group 1 - F column values (by row order) ===")
print(data1)
print("\n=== Group 2 - F column values (by row order) ===")
print(data2)

# Save results to a new Excel file (optional)
new_wb = Workbook()
ws1 = new_wb.active
ws1.title = "Group1_Row_F"
ws1.append(["Row", "F Value"])
for row, val in zip(target_rows1, data1):
    ws1.append([row, val])

ws2 = new_wb.create_sheet(title="Group2_Row_F")
ws2.append(["Row", "F Value"])
for row, val in zip(target_rows2, data2):
    ws2.append([row, val])

new_wb.save("extracted_f_values.xlsx")
print("\nDone. Saved: extracted_f_values.xlsx")
